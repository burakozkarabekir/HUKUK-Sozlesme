"""Excel export of an impact analysis (one row per affected clause)."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.impact import LEVEL_LABELS

DISCLAIMER = (
    "Bu çıktılar yapay zeka tarafından üretilmiştir ve hukuki görüş niteliği taşımaz. "
    "Nihai değerlendirme Hukuk birimine aittir."
)
_FILLS = {
    "ETKILENIR": PatternFill("solid", fgColor="F8D7DA"),
    "MUHTEMELEN_ETKILENIR": PatternFill("solid", fgColor="FFE5C2"),
    "ETKILENMEZ": PatternFill("solid", fgColor="E9ECEF"),
    "DEGERLENDIRILEMEDI": PatternFill("solid", fgColor="FFF3CD"),
}
_HEADER_FILL = PatternFill("solid", fgColor="1F3A68")
_COLUMNS = [
    ("Etki Seviyesi", 22),
    ("Güven", 9),
    ("Sözleşme No", 18),
    ("Başlık", 32),
    ("Taraflar", 36),
    ("Sözleşme Tipi", 22),
    ("Durum", 12),
    ("Madde", 18),
    ("Sayfa", 8),
    ("Ek Dosya", 20),
    ("Alıntı", 60),
    ("Alıntı Doğrulandı", 12),
    ("Gerekçe", 60),
    ("Önerilen Aksiyon", 40),
    ("Özet", 50),
]
# Spreadsheet apps execute cells starting with these characters as formulas.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _rows(results: list[dict[str, Any]]) -> list[list[Any]]:
    rows = []
    for result in results:
        base = [
            LEVEL_LABELS.get(result["impact_level"], result["impact_level"]),
            round(result["confidence"], 2) if result.get("confidence") is not None else None,
            result.get("contract_no") or result["contract_id"],
            result.get("title"),
            ", ".join(result.get("parties") or []),
            result.get("contract_type"),
            result.get("status"),
        ]
        tail = [result.get("recommended_action") or result.get("error"), result.get("summary")]
        clauses = result.get("affected_clauses") or []
        if not clauses:
            rows.append([*base, None, None, None, None, None, None, *tail])
        for clause in clauses:
            rows.append(
                [
                    *base,
                    clause.get("clause_ref"),
                    clause.get("page"),
                    clause.get("file_name") if clause.get("is_attachment") else None,
                    clause.get("quote"),
                    "Evet" if clause.get("quote_verified") else "HAYIR ⚠",
                    clause.get("reason"),
                    *tail,
                ]
            )
    return rows


def build_workbook(job: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sonuçlar"
    ws.append([name for name, _ in _COLUMNS])
    for index, (_, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    level_by_label = {label: level for level, label in LEVEL_LABELS.items()}
    for row in _rows(job.get("results", [])):
        ws.append([safe_cell(v) for v in row])
        fill = _FILLS.get(level_by_label.get(row[0], ""))
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            ws.cell(row=ws.max_row, column=1).fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{max(ws.max_row, 1)}"

    info = wb.create_sheet("Analiz Bilgisi")
    extraction = job.get("extraction") or {}
    filters = job.get("filters") or {}
    counts = job.get("counts") or {}
    lines = [
        ("Analiz", job.get("title")),
        ("Analiz No", job.get("id")),
        ("Oluşturan", job.get("created_by")),
        ("Oluşturma", job.get("created_at")),
        ("Durum", job.get("status")),
        ("Değerlendirilen sözleşme", f"{job.get('processed', 0)} / {job.get('total', 0)}"),
        *[(f"Sonuç: {LEVEL_LABELS[level]}", counts.get(level, 0)) for level in LEVEL_LABELS],
        ("Mevzuat özeti", extraction.get("summary")),
        ("Konular", ", ".join(extraction.get("topics") or [])),
        ("Anahtar kelimeler", ", ".join(job.get("keywords") or [])),
        ("Filtreler", ", ".join(f"{k}={v}" for k, v in filters.items() if v) or "-"),
        ("Rapor tarihi", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Uyarı", DISCLAIMER),
    ]
    for key, value in lines:
        info.append([key, safe_cell(value)])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
