from __future__ import annotations

import asyncio
import io
import json
import re
import time
from pathlib import Path

import openai
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import connect, migrate, utcnow_iso
from app.ingest.pipeline import run_sync
from app.llm.client import LLMClient
from app.main import create_app
from app.services.contracts import ContractFilters
from app.services.export import safe_cell
from app.services.impact import clean_keywords, find_candidates, verify_quote
from app.services.jobs import JobRunner
from tests import fakes
from tests.conftest import make_settings
from tests.pdfgen import make_pdf, write_sample_contracts

P = "/sozlesme"
LEGISLATION = (
    "7654 sayılı Kanun ile bayilik sözleşmelerinde fesih bildirim süresi en az 180 gün olarak belirlenmiştir. "
    "Gizlilik yükümlülükleri en fazla 3 yıl ile sınırlandırılmıştır."
)
EXTRACTION = {
    "summary": "Fesih bildirim süresi 180 güne çıkarıldı; gizlilik süresi 3 yılla sınırlandı.",
    "topics": ["fesih bildirim süresi", "gizlilik süresi"],
    "keywords": ["fesih", "feshedebilir", "bildirim süresi", "gizlilik", "gizli tutar", "süre"],
    "contract_types": ["Bayilik Sözleşmesi"],
    "effective_date": None,
}


def evaluation(contract_id: str, level: str, clauses: list[dict]) -> str:
    return json.dumps(
        {
            "contract_id": "MODEL-UYDURDU",  # the service must ignore this
            "impact_level": level,
            "confidence": 0.8 if level != "ETKILENMEZ" else 0.9,
            "affected_clauses": clauses,
            "recommended_action": "Madde tadil edilmeli" if level == "ETKILENIR" else "İşlem gerekmez",
            "summary": f"{contract_id} değerlendirmesi",
        },
        ensure_ascii=False,
    )


def responder(kwargs):
    instructions = kwargs["instructions"]
    text = kwargs["input"] if isinstance(kwargs["input"], str) else kwargs["input"][-1]["content"]
    if "kavramları çıkarmak" in instructions:
        return fakes.response(json.dumps(EXTRACTION, ensure_ascii=False))
    contract_id = re.search(r"contract_id\): (\S+)", text).group(1)
    if contract_id == "local-abc-bayilik":
        sources = re.findall(r'kaynak="(\d+)" madde="([^"]+)"', text)
        fesih = next(n for n, ref in sources if ref == "Madde 3")
        return fakes.response(
            evaluation(
                contract_id,
                "ETKILENIR",
                [
                    {
                        "source": int(fesih),
                        "clause_ref": "Madde 3",
                        "page": None,
                        "quote": "TARAFLARDAN her biri 90 gün önceden ... sözleşmeyi feshedebilir",
                        "reason": "Bildirim süresi 180 günün altında.",
                    },
                    {
                        "source": None,
                        "clause_ref": "Madde 4",
                        "page": 2,
                        "quote": "Bayi sırları ömür boyu saklar",
                        "reason": "Gizlilik süresi sınırı aşılıyor.",
                    },
                ],
            )
        )
    return fakes.response(evaluation(contract_id, "ETKILENMEZ", []))


@pytest.fixture
def env(tmp_path: Path):
    contracts = tmp_path / "contracts"
    write_sample_contracts(contracts)
    settings = make_settings(tmp_path, LOCAL_CONTRACTS_DIR=contracts, JOB_POLL_SECONDS=0.05)
    run_sync(settings)
    return settings


def client_for(settings, api) -> TestClient:
    app = create_app(settings)
    app.state.llm = LLMClient(settings, responses_api=api)
    return TestClient(app)


def wait_for(client: TestClient, job_id: str, timeout: float = 15) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"{P}/api/impact/{job_id}").json()
        if job["status"] in ("completed", "failed", "cancelled"):
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


# --- unit ---


def test_verify_quote() -> None:
    chunks = [
        {
            "chunk_id": 1,
            "text": "MADDE 3 - FESİH\n3.1 Taraflardan her biri 90 gün önceden yazılı\nbildirimde bulunarak sözleşmeyi feshedebilir.",
        },
        {"chunk_id": 2, "text": "Gizlilik 5 yıl sürer."},
    ]
    assert verify_quote("taraflardan HER BİRİ 90 gün önceden yazılı bildirimde", chunks)["chunk_id"] == 1
    assert verify_quote("“Taraflardan her biri …  sözleşmeyi feshedebilir.”", chunks)["chunk_id"] == 1
    assert verify_quote("sözleşmeyi feshedebilir ... taraflardan her biri", chunks) is None  # wrong order
    assert verify_quote("Taraflardan her biri 30 gün önceden", chunks) is None
    assert verify_quote("ab", chunks) is None
    assert verify_quote("Gizlilik 5 yıl sürer", chunks, preferred=2)["chunk_id"] == 2


def test_clean_keywords() -> None:
    assert clean_keywords([" Fesih ", "FESİH", '"sona erme"', "", "!!!", "x" * 81]) == ["Fesih", "sona erme"]


def test_find_candidates_groups_per_contract(env) -> None:
    conn = connect(env.db_path)
    migrate(conn)
    candidates = find_candidates(conn, ["fesih", "cezai şart"], ContractFilters(), per_contract=2, max_contracts=10)
    assert set(candidates) == {"local-abc-bayilik", "local-xyz-servis"}
    assert all(len(v) <= 2 for v in candidates.values())
    assert candidates["local-abc-bayilik"][0]["clause_ref"] == "Madde 3"
    only_one = find_candidates(conn, ["fesih", "cezai şart"], ContractFilters(), per_contract=2, max_contracts=1)
    assert len(only_one) == 1
    filtered = find_candidates(
        conn, ["fesih", "cezai şart"], ContractFilters(party="xyz"), per_contract=2, max_contracts=10
    )
    assert set(filtered) == {"local-xyz-servis"}


def test_safe_cell_blocks_formulas() -> None:
    assert safe_cell('=HYPERLINK("x")') == '\'=HYPERLINK("x")'
    assert safe_cell("-5 gün") == "'-5 gün" and safe_cell("normal") == "normal" and safe_cell(3) == 3


# --- API end to end ---


def test_impact_end_to_end(env) -> None:
    api = fakes.RuleResponses(responder)
    with client_for(env, api) as client:
        extracted = client.post(f"{P}/api/impact/extract", json={"text": LEGISLATION}).json()
        assert extracted["extraction"]["keywords"][:2] == ["fesih", "feshedebilir"]
        assert "<mevzuat_metni>" in api.calls[0]["input"]

        payload = {
            "title": "7654 sayılı Kanun",
            "text": extracted["text"],
            "keywords": extracted["extraction"]["keywords"],
            "extraction": extracted["extraction"],
            "filters": {"active_only": False},
        }
        started = client.post(f"{P}/api/impact/run", json=payload)
        assert started.status_code == 202
        job = wait_for(client, started.json()["job_id"])

        assert job["status"] == "completed" and job["total"] == job["processed"] == 2
        assert job["counts"] == {"ETKILENIR": 1, "ETKILENMEZ": 1}
        first, second = job["results"]
        assert first["contract_id"] == "local-abc-bayilik" and first["impact_level"] == "ETKILENIR"
        assert first["contract_no"] == "ABC_Bayilik" and second["impact_level"] == "ETKILENMEZ"
        verified, hallucinated = first["affected_clauses"]
        assert verified["quote_verified"] and verified["page"] == 2 and verified["clause_ref"] == "Madde 3"
        assert verified["chunk_id"] and verified["file_id"] == "abc-bayilik"
        assert not hallucinated["quote_verified"] and first["unverified_quotes"] == 1

        history = client.get(f"{P}/api/impact/history").json()["items"]
        assert history[0]["id"] == job["id"] and history[0]["title"] == "7654 sayılı Kanun"
        assert "results" not in history[0]

        xlsx = client.get(f"{P}/api/impact/{job['id']}/export.xlsx")
        assert xlsx.status_code == 200 and "attachment" in xlsx.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(xlsx.content))
        rows = list(wb["Sonuçlar"].iter_rows(values_only=True))
        assert rows[0][:3] == ("Etki Seviyesi", "Güven", "Sözleşme No")
        assert len(rows) == 4  # header + 2 clauses + 1 unaffected contract
        assert rows[1][0] == "Etkilenir" and rows[1][11] == "Evet" and rows[2][11].startswith("HAYIR")
        assert wb["Sonuçlar"].auto_filter.ref.startswith("A1:")
        info = dict(wb["Analiz Bilgisi"].iter_rows(values_only=True))
        assert "hukuki görüş niteliği taşımaz" in info["Uyarı"]

        conn = connect(env.db_path)
        actions = [r[0] for r in conn.execute("SELECT action FROM audit_log ORDER BY id")]
        assert {"impact.extract", "impact.run", "impact.export"} <= set(actions)
        details = " ".join(r[0] for r in conn.execute("SELECT details_json FROM audit_log"))
        assert "180 gün" not in details  # legislation text is not audited


def test_extract_from_pdf_and_limits(env) -> None:
    api = fakes.RuleResponses(responder)
    settings = env.model_copy(update={"MAX_UPLOAD_MB": 1, "MAX_LEGISLATION_CHARS": 50})
    with client_for(settings, api) as client:
        pdf = make_pdf(["Kanun metni: fesih bildirim süresi.\nİkinci satır."])
        response = client.post(f"{P}/api/impact/extract", content=pdf, headers={"Content-Type": "application/pdf"})
        assert response.status_code == 200 and response.json()["input_kind"] == "pdf"
        assert "fesih" in response.json()["text"]
        assert (
            client.post(
                f"{P}/api/impact/extract", content=b"not pdf", headers={"Content-Type": "application/pdf"}
            ).status_code
            == 415
        )
        scanned = client.post(
            f"{P}/api/impact/extract", content=make_pdf([""]), headers={"Content-Type": "application/pdf"}
        )
        assert scanned.status_code == 422 and scanned.json()["error"]["code"] == "pdf_needs_ocr"
        big = client.post(
            f"{P}/api/impact/extract",
            content=b"%PDF" + b"0" * (1024 * 1024 + 1),
            headers={"Content-Type": "application/pdf"},
        )
        assert big.status_code == 413
        too_long = client.post(f"{P}/api/impact/extract", json={"text": "x" * 51})
        assert too_long.status_code == 413 and too_long.json()["error"]["code"] == "legislation_too_long"
        assert (
            client.post(f"{P}/api/impact/extract", content=b"a", headers={"Content-Type": "text/plain"}).status_code
            == 415
        )


def test_fatal_llm_error_fails_job(env) -> None:
    api = fakes.RuleResponses(lambda kw: fakes.api_error(openai.AuthenticationError, 401, "bad key"))
    with client_for(env, api) as client:
        job_id = client.post(f"{P}/api/impact/run", json={"text": LEGISLATION, "keywords": ["fesih", "cezai"]}).json()[
            "job_id"
        ]
        job = wait_for(client, job_id)
    assert job["status"] == "failed" and "API anahtarını" in job["error_message"]


def test_per_contract_error_is_reported_not_fatal(env) -> None:
    def flaky(kwargs):
        if "local-xyz-servis" in kwargs["input"]:
            return fakes.api_error(openai.InternalServerError, 500, "boom")
        return responder(kwargs)

    with client_for(env, fakes.RuleResponses(flaky)) as client:
        job_id = client.post(f"{P}/api/impact/run", json={"text": LEGISLATION, "keywords": ["fesih", "cezai"]}).json()[
            "job_id"
        ]
        job = wait_for(client, job_id)
    assert job["status"] == "completed"
    failed = next(r for r in job["results"] if r["contract_id"] == "local-xyz-servis")
    assert failed["impact_level"] == "DEGERLENDIRILEMEDI" and "geçici bir hata" in failed["error"]


def test_cancel(env) -> None:
    gate = asyncio.Event()

    class SlowApi:
        calls = 0

        async def create(self, **kwargs):
            SlowApi.calls += 1
            await asyncio.wait_for(gate.wait(), timeout=10)
            return responder(kwargs)

    app = create_app(env)
    app.state.llm = LLMClient(env, responses_api=SlowApi())
    with TestClient(app) as client:
        job_id = client.post(f"{P}/api/impact/run", json={"text": LEGISLATION, "keywords": ["fesih", "cezai"]}).json()[
            "job_id"
        ]
        deadline = time.time() + 5
        while SlowApi.calls == 0 and time.time() < deadline:
            time.sleep(0.02)
        cancelled = client.post(f"{P}/api/impact/{job_id}/cancel").json()
        assert cancelled["cancel_requested"]
        job = wait_for(client, job_id)
        assert job["status"] == "cancelled"
        assert client.post(f"{P}/api/impact/{job_id}/cancel").status_code == 409
    assert client.get  # silence linters


def test_cancel_queued_job_and_unknown(env) -> None:
    conn = connect(env.db_path)
    migrate(conn)
    from app.services.impact import RunRequest, create_job, request_cancel

    job_id = create_job(conn, RunRequest(text="x", keywords=["fesih"]), "ali")
    assert request_cancel(conn, job_id)["status"] == "cancelled"
    with TestClient(create_app(env)) as client:
        assert client.get(f"{P}/api/impact/yok").status_code == 404
        assert client.post(f"{P}/api/impact/run", json={"text": "x", "keywords": []}).status_code == 422


def test_stale_running_job_is_requeued_and_completed(env) -> None:
    conn = connect(env.db_path)
    migrate(conn)
    from app.services.impact import RunRequest, create_job

    job_id = create_job(conn, RunRequest(text=LEGISLATION, keywords=["fesih"]), "ali")
    conn.execute(
        "UPDATE analysis_jobs SET status='running', worker_id='dead', attempts=1, heartbeat_at='2000-01-01T00:00:00+00:00' WHERE id=?",
        (job_id,),
    )

    async def scenario() -> str:
        runner = JobRunner(env, LLMClient(env, responses_api=fakes.RuleResponses(responder)))
        runner.start()
        try:
            for _ in range(200):
                status = conn.execute("SELECT status FROM analysis_jobs WHERE id=?", (job_id,)).fetchone()[0]
                if status == "completed":
                    return status
                await asyncio.sleep(0.05)
            return status
        finally:
            await runner.stop()

    assert asyncio.run(scenario()) == "completed"
    assert conn.execute("SELECT attempts FROM analysis_jobs WHERE id=?", (job_id,)).fetchone()[0] == 2


def test_job_claimed_once_by_two_runners(env) -> None:
    conn = connect(env.db_path)
    migrate(conn)
    from app.services.impact import RunRequest, create_job

    job_id = create_job(conn, RunRequest(text=LEGISLATION, keywords=["fesih"]), "ali")
    first = JobRunner(env, LLMClient(env, responses_api=fakes.RuleResponses(responder)))
    second = JobRunner(env, LLMClient(env, responses_api=fakes.RuleResponses(responder)))
    assert first._claim() == job_id
    assert second._claim() is None
    assert conn.execute("SELECT status, worker_id FROM analysis_jobs").fetchone()[1] == first.worker_id
    _ = utcnow_iso()
